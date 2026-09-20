"""Text-layer extraction with stable locators. No OCR, macros, scripts or links run."""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from html.parser import HTMLParser
from pathlib import Path

from .common import ReviewError, digest, load_json

EXTRACTOR_VERSION = "1.0.0"
SUPPORTED = {".txt", ".md", ".markdown", ".html", ".htm", ".json", ".pdf", ".docx", ".pptx"}


@dataclass
class Segment:
    id: str
    text: str
    locator: str
    start: int
    end: int
    excluded: bool = False
    exclusion_reason: str | None = None


class VisibleHTML(HTMLParser):
    """Conservative text extraction, not a browser's computed visibility model."""

    BLOCK = {"p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "li", "br", "tr", "section", "article", "table"}
    SKIP = {"script", "style", "noscript", "template", "head", "svg", "canvas"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.stack = []

    def handle_starttag(self, tag, attrs):
        attributes = dict(attrs)
        hidden = tag in self.SKIP or "hidden" in attributes or attributes.get("aria-hidden") == "true"
        hidden = hidden or bool(
            re.search(r"display\s*:\s*none|visibility\s*:\s*hidden", attributes.get("style", ""), re.I)
        )
        inherited = any(self.stack)
        if tag not in {"br", "img", "hr", "meta", "link", "input", "source", "wbr"}:
            self.stack.append(hidden or inherited)
        if tag in self.BLOCK and not hidden and not inherited:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if self.stack:
            self.stack.pop()
        if tag in self.BLOCK and not any(self.stack):
            self.parts.append("\n")

    def handle_data(self, data):
        if not any(self.stack):
            self.parts.append(data)


def _read_source(path: Path) -> tuple[list[tuple[str, str]], list[str], str]:
    """Return text blocks, explicit coverage warnings, and extraction scope."""
    suffix = path.suffix.lower()
    warnings = []
    if suffix not in SUPPORTED:
        raise ReviewError(f"Unsupported format {suffix}. Export prose to UTF-8 Markdown or JSON.")
    if path.stat().st_size > 30_000_000:
        raise ReviewError("Input exceeds the 30 MB safety limit; split or export the document.")
    if suffix in {".md", ".markdown", ".txt"}:
        return [("document", path.read_text(encoding="utf-8-sig"))], warnings, "source_text"
    if suffix in {".html", ".htm"}:
        parser = VisibleHTML()
        parser.feed(path.read_text(encoding="utf-8-sig"))
        parser.close()
        warnings.append("HTML text only: CSS-computed visibility, charts, images and layout are not assessed.")
        return [("html", "".join(parser.parts))], warnings, "html_visible_text_approximation"
    if suffix == ".json":
        obj = load_json(path)
        # An explicit prose export, not arbitrary serialized business data.
        if isinstance(obj, dict) and isinstance(obj.get("text"), str):
            return [("document", obj["text"])], warnings, "exported_text"
        if isinstance(obj, dict) and isinstance(obj.get("blocks"), list):
            out = []
            for index, b in enumerate(obj["blocks"], 1):
                if not isinstance(b, dict) or not isinstance(b.get("text"), str):
                    raise ReviewError("Each JSON block must contain a text string.")
                out.append((str(b.get("locator", f"block {index}")), b["text"]))
            return out, warnings, "exported_blocks"
        raise ReviewError('JSON input must be {"text":"..."} or {"blocks":[{"locator":"...","text":"..."}]}.')
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ReviewError("Install optional readers: pip install -r requirements.txt") from exc
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ReviewError("Encrypted PDF: provide an authorized unencrypted text export.")
        out = []
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if len(text.strip()) < 10:
                warnings.append(f"PDF page {i}: no usable text layer. This page is not evaluated; no OCR was run.")
                text = ""
            out.append((f"page {i}", text))
        warnings.append(
            "PDF text layer only: reading order may be imperfect; images, visual design and chart values are not assessed."
        )
        return out, warnings, "pdf_text_layer"
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ReviewError("Install optional readers: pip install -r requirements.txt") from exc
        document = Document(str(path))
        out = []
        # Iterating OOXML body preserves ordinary paragraph/table order.
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        for i, child in enumerate(document.element.body, 1):
            paragraphs = list(child.iter(ns + "p"))
            if child.tag == ns + "p":
                paragraphs = [child]
            for j, p in enumerate(paragraphs, 1):
                text = "".join(n.text or "" for n in p.iter(ns + "t"))
                if text:
                    out.append((f"body block {i}, paragraph {j}", text))
        warnings.append(
            "DOCX body text only: headers, footers, comments, tracked-change semantics, images and visual layout are not assessed."
        )
        return out, warnings, "docx_body_text"
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ReviewError("Install optional readers: pip install -r requirements.txt") from exc
    presentation = Presentation(str(path))
    out = []

    def visit(shapes, prefix):
        for j, shape in enumerate(shapes, 1):
            location = f"{prefix}, shape {j}"
            if hasattr(shape, "shapes"):
                visit(shape.shapes, location)
            elif shape.has_text_frame:
                out.append((location, shape.text_frame.text))
            elif shape.has_table:
                for k, row in enumerate(shape.table.rows, 1):
                    out.append((f"{location}, row {k}", " | ".join(c.text for c in row.cells)))

    for i, slide in enumerate(presentation.slides, 1):
        visit(slide.shapes, f"slide {i}")
    warnings.append(
        "PPTX shape/table text only: images, charts, notes, slide ordering within groups and visual design are not assessed."
    )
    return out, warnings, "pptx_text"


def language_of(text: str) -> str:
    japanese = len(re.findall(r"[\u3040-\u30ff\u3400-\u9fff]", text))
    latin = len(re.findall(r"[A-Za-z]", text))
    if japanese and latin > 2 * japanese:
        return "mixed"
    if japanese:
        return "ja"
    if latin:
        return "en"
    return "unknown"


def extract(path: Path, locked_sections: list[str] | None = None) -> dict:
    try:
        blocks, warnings, scope = _read_source(path)
    except (OSError, UnicodeError) as exc:
        raise ReviewError(f"Cannot read file: {path.name}; UTF-8 text is required for text inputs.") from exc
    normalized = ""
    segments = []
    fence = None
    locked_sections = locked_sections or []
    for loc, text in blocks:
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        if normalized:
            normalized += "\n\n"
        block_start = len(normalized)
        normalized += text
        offset = 0
        for line_number, line in enumerate(text.splitlines(keepends=True), 1):
            body = line.rstrip("\n")
            stripped = body.strip()
            marker = re.match(r"^\s*(`{3,}|~{3,})", body)
            in_fence = fence is not None
            if marker:
                token = marker.group(1)[0]
                if fence is None:
                    fence = token
                elif fence == token:
                    fence = None
                in_fence = True
            quote = bool(re.match(r"^\s*>", body))
            locked = bool(stripped and any(stripped in item or item in stripped for item in locked_sections))
            reason = "code_fence" if in_fence else ("blockquote" if quote else ("locked_section" if locked else None))
            # Bound spans without manufacturing text; every quote is an exact normalized-source slice.
            for part in re.finditer(r"[^。！？!?\n]+[。！？!?]?|[。！？!?]", body):
                pstart = part.start()
                pend = part.end()
                for base in range(pstart, pend, 700):
                    end = min(base + 700, pend)
                    raw = body[base:end]
                    leading = len(raw) - len(raw.lstrip())
                    trailing = len(raw.rstrip())
                    begin = base + leading
                    finish = base + trailing
                    if finish <= begin:
                        continue
                    start = block_start + offset + begin
                    stop = block_start + offset + finish
                    segments.append(
                        Segment(
                            f"s{len(segments) + 1:05}",
                            normalized[start:stop],
                            f"{loc}, line {line_number}",
                            start,
                            stop,
                            bool(reason),
                            reason,
                        )
                    )
            offset += len(line)
    if not normalized.strip():
        raise ReviewError("No extractable text. Supply a text export; scanned images were not evaluated.")
    if len(normalized) > 500_000:
        raise ReviewError("Extracted text exceeds 500,000 characters; split by document/section.")
    return {
        "name": path.name,
        "text": normalized,
        "sha256": digest(normalized),
        "scope": scope,
        "warnings": warnings,
        "segments": [asdict(s) for s in segments],
        "detected_language": language_of(normalized),
        "extractor_version": EXTRACTOR_VERSION,
    }


def make_chunks(document: dict, max_chars: int = 2600, max_chunks: int = 128) -> list[dict]:
    chunks = []
    current = []
    length = 0
    for seg in document["segments"]:
        if seg["excluded"]:
            continue
        if current and (length + len(seg["text"]) > max_chars or len(current) >= 32):
            chunks.append({"id": f"c{len(chunks) + 1:04}", "segments": current, "chars": length})
            current = []
            length = 0
        current.append(seg)
        length += len(seg["text"])
    if current:
        chunks.append({"id": f"c{len(chunks) + 1:04}", "segments": current, "chars": length})
    if len(chunks) > max_chunks:
        raise ReviewError(f"{len(chunks)} chunks exceed --max-chunks {max_chunks}; no partial score was produced.")
    return chunks
